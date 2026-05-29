# app.py
import streamlit as st
from mistralai import Mistral
from dotenv import load_dotenv
from docx import Document
from docxcompose.composer import Composer
from PyPDF2 import PdfReader
import os

# Load Environment
load_dotenv()

# Mistral Client
client = Mistral(
    api_key=os.getenv("MISTRAL_API_KEY")
)

# Page Config
st.set_page_config(
    page_title="AI PDF Answer Generator",
    page_icon="📘",
    layout="wide"
)

# Title
st.title("📘 AI PDF Answer Generator")
st.write("Upload PDF → Extract Questions → Generate AI Answers → Merge DOCX")

# Upload PDF
uploaded_file = st.file_uploader(
    "Upload PDF File",
    type=["pdf"]
)

# Generate Button
if uploaded_file is not None:

    if st.button("Generate Answers"):

        # Save Uploaded File
        with open("temp.pdf", "wb") as f:
            f.write(uploaded_file.read())

        st.success("PDF Uploaded Successfully!")

        # Read PDF
        reader = PdfReader("temp.pdf")

        questions = []

        # Extract Questions
        for page in reader.pages:

            text = page.extract_text()

            if text:

                lines = text.split("\n")

                for line in lines:

                    line = line.strip()

                    if "." in line:

                        first_part = line.split(".", 1)[0]

                        if first_part.isdigit():

                            q = line.split(".", 1)[1].strip()

                            questions.append(q)

        st.write(f"✅ Total Questions Found: {len(questions)}")

        # File List
        file_names = []

        # Progress Bar
        progress = st.progress(0)

        # Generate Answers
        for i, pro in enumerate(questions):

            prompt = f"""
            Question: {pro}

            Write a clear, well-structured answer with:
            - 10 key points
            - examples
            - conclusion
            Use simple English.
            """

            response = client.chat.complete(
                model="mistral-small-latest",
                messages=[
                    {
                        "role": "user",
                        "content": prompt
                    }
                ]
            )

            opt = response.choices[0].message.content

            # Create DOCX
            doc = Document()

            doc.add_heading(f"Question {i+1}", level=1)
            doc.add_paragraph(pro)

            doc.add_heading("Answer", level=2)
            doc.add_paragraph(opt)

            filename = f"output{i+1}.docx"

            doc.save(filename)

            file_names.append(filename)

            progress.progress((i + 1) / len(questions))

        st.success("All Answers Generated!")

        # Merge DOCX Files
        if len(file_names) > 0:

            main_doc = Document(file_names[0])

            composer = Composer(main_doc)

            for file in file_names[1:]:

                temp_doc = Document(file)

                composer.append(temp_doc)

            composer.save("merged_output.docx")

            st.success("DOCX Files Merged Successfully!")

            # Download Button
            with open("merged_output.docx", "rb") as file:

                st.download_button(
                    label="📥 Download Merged DOCX",
                    data=file,
                    file_name="merged_output.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )